import pandas as pd
import openpyxl
from openpyxl.utils import get_column_letter
import os
import glob
import logging
from datetime import datetime

def excel_col_to_number(col_letter):
    """Excel列のアルファベットを列番号に変換（A=1, B=2, ..., Z=26, AA=27, ...）"""
    col = 0
    for char in col_letter:
        col = col * 26 + (ord(char) - ord('A') + 1)
    return col

def get_delete_ranges(filename):
    """
    ファイル名からキーワードを検出し、対応する削除列範囲を返す
    
    Args:
        filename: ファイル名
    
    Returns:
        削除する列の範囲リスト [('A', 'BV'), ('EY', 'GB'), ...]
    """
    delete_config = {
        'ASSY 試作用1': [('A', 'BV'), ('EX', 'GA')],
        'ASSY1': [('A', 'BV'), ('EY', 'GB')],
        'CNC11': [('A', 'BV'), ('EH', 'EN'), ('EP', 'ES'), ('EU', 'FX'), ('JU', 'KS'), ('TA', 'TW'), ('TY', 'TY')],
        'CNC21': [('A', 'BV'), ('EH', 'EN'), ('EP', 'ES'), ('EU', 'FX'), ('JU', 'KS'), ('TA', 'TW'), ('TY', 'TY')],
        'S1研磨1': [('A', 'BV'), ('EN', 'ET'), ('EV', 'GA'), ('OR', 'PM')],
        'S1切断1': [('A', 'BV'), ('EH', 'EO'), ('EQ', 'FU'), ('JR', 'KP'), ('SX', 'TT')],
        'S2研磨1': [('A', 'BV'), ('EN', 'ET'), ('EV', 'GA'), ('OR', 'PM')],
        'S2切断1': [('A', 'BV'), ('EH', 'EO'), ('EQ', 'FU'), ('JR', 'KP'), ('SX', 'TT')],
        'マーク1': [('A', 'BV'), ('DA', 'DD')],
        'ラミ1': [('A', 'BV'), ('EM', 'FQ'), ('GC', 'GC')],
        '印刷 タブレット用': [('A', 'BV'), ('LR', 'MQ')],
        '印刷': [('A', 'BV'), ('ET', 'EX')],
        '強化1': [('A', 'BV'), ('FB', 'FF')],
        '研磨1': [('A', 'BV'), ('EM', 'ER'), ('ET', 'ET'), ('EV', 'GA'), ('OR', 'PM')],
        '孔あけ1': [('A', 'BV'), ('EH', 'FX'), ('JK', 'JK')],
        '切断＋研磨1': [('A', 'BV'), ('EN', 'GD'), ('KA', 'KZ'), ('OW', 'PV'), ('YD', 'ZA')],
        '切断1': [('A', 'BV'), ('EH', 'EN'), ('EP', 'ES'), ('EU', 'FX'), ('JU', 'KS'), ('TA', 'TW'), ('TY', 'TY')],
        '特殊作業1': [('A', 'BV'), ('EP', 'ET')],
        '品証1': [('A', 'BV'), ('EH', 'EL')],
    }
    
    # ファイル名からキーワードを検出（長いキーワードから順にチェック）
    for keyword in sorted(delete_config.keys(), key=len, reverse=True):
        if keyword in filename:
            return delete_config[keyword], keyword
    
    return None, None

def delete_columns_from_csv(input_file, output_file, delete_ranges, keyword=''):
    """
    CSVファイルから指定した列を削除する
    
    Args:
        input_file: 入力CSVファイルパス
        output_file: 出力CSVファイルパス
        delete_ranges: 削除する列の範囲リスト [('A', 'BV'), ('EY', 'GB'), ...]
        keyword: ファイルを識別するキーワード
    """
    for enc in ['cp932','shift_jis','utf-8','utf-8-sig','latin1']:
        try:
            df = pd.read_csv(input_file, encoding=enc)
            print('loaded with', enc)
            break
        except Exception:
            continue
    else:
       raise RuntimeError('読み込みに全て失敗しました')
    try:
        # CSVファイルを読み込む（文字化け対策）
        # df = pd.read_csv(input_file, encoding='enc')
        print(f"  元のデータ形状: {df.shape}")
        print(f"  元の列数: {len(df.columns)}")
        
        # 削除する列番号を収集
        cols_to_delete = []
        
        for start_col, end_col in delete_ranges:
            start_num = excel_col_to_number(start_col)
            end_num = excel_col_to_number(end_col)
            
            print(f"  削除範囲: {start_col}({start_num}) ～ {end_col}({end_num})")
            
            # 列インデックスは0ベースなので、-1する
            for col_num in range(start_num - 1, end_num):
                if col_num < len(df.columns):
                    cols_to_delete.append(col_num)
        
        # 重複を削除してソート
        cols_to_delete = sorted(set(cols_to_delete))
        print(f"  削除する列数: {len(cols_to_delete)}")
        
        # 削除する列のインデックスで列を削除
        df_result = df.drop(df.columns[cols_to_delete], axis=1)
        
        print(f"  削除後のデータ形状: {df_result.shape}")
        print(f"  削除後の列数: {len(df_result.columns)}")
        
        # S2研磨1の特殊処理：69列目と70列目を入れ替える
        if keyword == 'S2研磨1':
            try:
                if len(df_result.columns) >= 70:
                    print(f"  特殊処理: Col 69（設備コード）とCol 70（受注番号）を入れ替えます")
                    # 69列目と70列目のインデックス（0ベース）
                    col69_idx = 68
                    col70_idx = 69
                    
                    # 列を入れ替え
                    cols = list(df_result.columns)
                    cols[col69_idx], cols[col70_idx] = cols[col70_idx], cols[col69_idx]
                    df_result = df_result[cols]
                    print(f"  ✓ 列の入れ替えが完了しました")
            except Exception as e:
                print(f"  ⚠ 列の入れ替え処理中にエラーが発生しました: {e}")
        
        # 結果をCSVに保存（UTF-8エンコーディング）
        df_result.to_csv(output_file, index=False, encoding='utf-8-sig')
        print(f"  ✓ ファイルを保存しました: {output_file}")
        return True
        
    except Exception as e:
        print(f"  ✗ エラーが発生しました: {e}")
        return False

def delete_columns_from_excel(input_file, output_file, delete_ranges, keyword=''):
    """
    Excelファイルから指定した列を削除する
    
    Args:
        input_file: 入力Excelファイルパス
        output_file: 出力Excelファイルパス
        delete_ranges: 削除する列の範囲リスト [('A', 'BV'), ('EY', 'GB'), ...]
        keyword: ファイルを識別するキーワード
    """
    try:
        # Excelファイルを読み込む
        wb = openpyxl.load_workbook(input_file)
        ws = wb.active
        
        print(f"  シート名: {ws.title}")
        print(f"  元の最大列数: {ws.max_column}")
        
        # 削除する列番号を逆順で収集（後ろから削除する）
        cols_to_delete = []
        
        for start_col, end_col in delete_ranges:
            start_num = excel_col_to_number(start_col)
            end_num = excel_col_to_number(end_col)
            
            print(f"  削除範囲: {start_col}({start_num}) ～ {end_col}({end_num})")
            
            for col_num in range(start_num, end_num + 1):
                cols_to_delete.append(col_num)
        
        # 重複を削除して降順でソート（後ろから削除するため）
        cols_to_delete = sorted(set(cols_to_delete), reverse=True)
        print(f"  削除する列数: {len(cols_to_delete)}")
        
        # 列を削除（後ろから削除することで列番号のずれを防ぐ）
        for col_num in cols_to_delete:
            ws.delete_cols(col_num)
        
        # S2研磨1の特殊処理：69列目（BQ）と70列目（BR）を入れ替える
        if keyword == 'S2研磨1':
            try:
                bq_col_num = 69  # BQ列
                br_col_num = 70  # BR列
                
                # 削除後に69列目と70列目が存在する場合のみ処理
                if bq_col_num <= ws.max_column and br_col_num <= ws.max_column:
                    print(f"  特殊処理: Col 69（BQ）とCol 70（BR）を入れ替えます")
                    
                    # BQ列のデータを一時保存
                    bq_data = []
                    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=bq_col_num, max_col=bq_col_num):
                        bq_data.append([cell.value for cell in row])
                    
                    # BR列のデータを一時保存
                    br_data = []
                    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=br_col_num, max_col=br_col_num):
                        br_data.append([cell.value for cell in row])
                    
                    # BQ列の位置にBR列のデータを配置
                    for row_idx, row_data in enumerate(br_data, 1):
                        ws.cell(row=row_idx, column=bq_col_num).value = row_data[0]
                    
                    # BR列の位置にBQ列のデータを配置
                    for row_idx, row_data in enumerate(bq_data, 1):
                        ws.cell(row=row_idx, column=br_col_num).value = row_data[0]
                    
                    print(f"  ✓ 列の入れ替えが完了しました")
            except Exception as e:
                print(f"  ⚠ 列の入れ替え処理中にエラーが発生しました: {e}")
        
        print(f"  削除後の最大列数: {ws.max_column}")
        
        # 結果を保存
        wb.save(output_file)
        print(f"  ✓ ファイルを保存しました: {output_file}")
        return True
        
    except Exception as e:
        print(f"  ✗ エラーが発生しました: {e}")
        return False

def process_all_files(directory='.'):
    """
    ディレクトリ内のすべてのCSV/Excelファイルを処理する
    
    Args:
        directory: 処理対象のディレクトリ
    """
    print("=" * 80)
    print("列削除自動処理プログラム")
    print("=" * 80)
    
    # CSVファイルを処理
    csv_files = glob.glob(os.path.join(directory, '*.csv'))
    excel_files = glob.glob(os.path.join(directory, '*.xlsx')) + glob.glob(os.path.join(directory, '*.xls'))
    
    all_files = csv_files + excel_files
    
    output_dir = os.path.join(directory, '削除済')
    if not os.path.exists(output_dir):
        os.makedirs(output_dir)

    # ログ設定（削除済フォルダにログを保存）
    #log_filename = f"処理ログ_{datetime.now().strftime('%Y%m%d_%H%M%S')}.txt"
    #log_path = os.path.join(output_dir, log_filename)
    #logging.basicConfig(filename=log_path, level=logging.INFO, format='%(asctime)s %(levelname)s: %(message)s')
    #logging.getLogger('delete_columns').info(f"処理開始: {datetime.now().isoformat()}")

    if not all_files:
        print("対象ファイルが見つかりません")
        logging.getLogger('delete_columns').info("対象ファイルが見つかりません")
        return
    
    success_count = 0
    fail_count = 0
    
    for input_file in sorted(all_files):
        filename = os.path.basename(input_file)
        delete_ranges, keyword = get_delete_ranges(filename)
        
        if delete_ranges is None:
            print(f"\n⊘ スキップ: {filename}")
            print(f"  対応するキーワードが見つかりません")
            continue
        
        # 出力ファイル名を生成（削除済フォルダへ、ファイル名は変更しない）
        output_file = os.path.join(output_dir, filename)
        
        print(f"\n▶ 処理中: {filename}")
        print(f"  キーワード: {keyword}")
        
        # ファイル形式に応じて処理
        if input_file.endswith('.csv'):
            if delete_columns_from_csv(input_file, output_file, delete_ranges, keyword):
                success_count += 1
            else:
                fail_count += 1
        else:
            if delete_columns_from_excel(input_file, output_file, delete_ranges, keyword):
                success_count += 1
            else:
                fail_count += 1
    
    print("\n" + "=" * 80)
    print(f"処理完了: 成功 {success_count}件, 失敗 {fail_count}件")
    print("=" * 80)

# メイン処理
if __name__ == "__main__":
    # カレントディレクトリを処理
    process_all_files()
